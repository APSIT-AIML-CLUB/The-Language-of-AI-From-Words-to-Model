"""
ResumeIQ AI — Unified Agentic Pipeline (Entry Point)
"""
import os
import json
from pathlib import Path
from typing import Any, Dict, Optional
from uuid import uuid4
from datetime import datetime
from contextlib import asynccontextmanager

import fitz  # PyMuPDF
import docx  # python-docx
from pydantic import BaseModel, Field, model_validator
from fastapi import FastAPI, UploadFile, File, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
import spacy
from sentence_transformers import SentenceTransformer

# Load environment variables
from dotenv import load_dotenv
load_dotenv()

import importlib

# Because the filenames start with numbers, standard import syntax (from 2_nlp import ...) 
# is syntactically invalid in Python. We use importlib to dynamically import them instead.
nlp_stage = importlib.import_module("2_nlp")
extract_skills = nlp_stage.extract_skills
parse_resume_text = nlp_stage.parse_resume_text
evaluate_student_readiness = nlp_stage.evaluate_student_readiness

gemini_stage = importlib.import_module("3_gemini")
generate_copilot_bundle = gemini_stage.generate_copilot_bundle

PROJECT_NAME = "ResumeIQ AI (Teaching Edition)"
VERSION = "2.0.0-teach"

# Load the spaCy model used for NLP processing
SPACY_MODEL = os.getenv("SPACY_MODEL", "FIXME").strip('"\'')
# Load the embedding model used for semantic comparison
EMBEDDING_MODEL = os.getenv("EMBEDDING_MODEL", "FIXME").strip('"\'')

ALLOWED_EXTENSIONS = [".pdf", ".docx"]

# ============================================================
# ROLE KNOWLEDGE BASE
# ============================================================
ROLE_REQUIREMENTS: Dict[str, Dict[str, Any]] = {
    "software engineer": {
        "title": "Software Engineer",
        "core_skills": ["Python", "Java", "Data Structures", "Algorithms", "Git", "REST APIs", "SQL"],
        "secondary_skills": ["Docker", "Linux", "CI/CD", "PostgreSQL", "FastAPI", "System Design"],
        "description": "Software engineers design, build, and maintain software systems: writing clean code, "
                        "building RESTful APIs, optimizing algorithms, working with relational databases, "
                        "and version controlling with Git.",
    },
    "frontend developer": {
        "title": "Frontend Developer",
        "core_skills": ["JavaScript", "TypeScript", "React", "HTML", "CSS", "Git", "REST APIs"],
        "secondary_skills": ["Next.js", "Tailwind CSS", "Vue", "Angular", "Redux", "GraphQL"],
        "description": "Frontend developers build responsive, dynamic, and intuitive user interfaces. "
                        "They translate UX wireframes into functional web applications using modern "
                        "JavaScript/TypeScript frameworks such as React and Next.js.",
    },
    "backend developer": {
        "title": "Backend Developer",
        "core_skills": ["Python", "FastAPI", "Node.js", "SQL", "PostgreSQL", "REST APIs", "Git"],
        "secondary_skills": ["Docker", "Redis", "Microservices", "MongoDB", "Linux", "AWS", "CI/CD"],
        "description": "Backend developers architect server-side logic, database schemas, and REST APIs. "
                        "They optimize system performance, ensure secure data management, and implement "
                        "asynchronous workflows.",
    },
    "full stack developer": {
        "title": "Full Stack Developer",
        "core_skills": ["JavaScript", "TypeScript", "React", "Node.js", "SQL", "REST APIs", "Git"],
        "secondary_skills": ["Next.js", "Python", "FastAPI", "Docker", "PostgreSQL", "MongoDB", "Tailwind CSS"],
        "description": "Full stack developers manage both client-facing interfaces and backend infrastructure. "
                        "They build end-to-end features from database queries to interactive web components.",
    },
    "machine learning engineer": {
        "title": "Machine Learning Engineer",
        "core_skills": ["Python", "Machine Learning", "Scikit-Learn", "Pandas", "NumPy", "Statistics", "Git"],
        "secondary_skills": ["PyTorch", "TensorFlow", "Deep Learning", "Natural Language Processing", "Docker", "SQL", "FastAPI", "LLMs"],
        "description": "Machine learning engineers research, build, and deploy predictive models. They perform "
                        "data preprocessing, feature engineering, model training, and integration into production APIs.",
    },
    "data analyst": {
        "title": "Data Analyst",
        "core_skills": ["SQL", "Python", "Excel", "Pandas", "Statistics", "Power BI"],
        "secondary_skills": ["Tableau", "NumPy", "Git", "PostgreSQL", "MySQL", "Data Visualization"],
        "description": "Data analysts inspect, cleanse, transform, and model data to discover useful information, "
                        "report findings, and support data-driven business decisions.",
    },
    "data scientist": {
        "title": "Data Scientist",
        "core_skills": ["Python", "Statistics", "Machine Learning", "Pandas", "NumPy", "SQL", "Git"],
        "secondary_skills": ["Scikit-Learn", "Deep Learning", "Natural Language Processing", "Data Visualization", "Docker", "AWS"],
        "description": "Data scientists combine domain expertise, programming skills, and mathematics to extract "
                        "meaningful insights and predictive signals from complex datasets.",
    },
    "devops engineer": {
        "title": "DevOps Engineer",
        "core_skills": ["Linux", "Docker", "Git", "CI/CD", "AWS", "Kubernetes"],
        "secondary_skills": ["Python", "Bash", "Terraform", "GCP", "Azure", "Microservices", "Redis"],
        "description": "DevOps engineers bridge software development and operations teams. They automate CI/CD "
                        "pipelines, orchestrate containerized environments with Docker and Kubernetes, and manage "
                        "cloud infrastructure.",
    },
    "cloud engineer": {
        "title": "Cloud Engineer",
        "core_skills": ["AWS", "Linux", "Docker", "Python", "Git", "SQL"],
        "secondary_skills": ["Azure", "GCP", "Kubernetes", "CI/CD", "Terraform", "Microservices"],
        "description": "Cloud engineers design, configure, and maintain cloud architectures. They migrate "
                        "workloads, implement serverless applications, and guarantee cloud security and scalability.",
    },
    "mobile app developer": {
        "title": "Mobile App Developer",
        "core_skills": ["Flutter", "React Native", "JavaScript", "TypeScript", "Git", "REST APIs"],
        "secondary_skills": ["Android", "iOS", "Swift", "Kotlin", "Firebase", "SQL"],
        "description": "Mobile developers create responsive, high-performance applications for iOS and Android, "
                        "integrating native APIs, local storage, and real-time backend services.",
    },
    "cyber security analyst": {
        "title": "Cyber Security Analyst",
        "core_skills": ["Linux", "Python", "Git", "SQL", "REST APIs"],
        "secondary_skills": ["Docker", "AWS", "Bash", "System Design"],
        "description": "Cyber security analysts protect systems, networks, and data from cyber attacks. They "
                        "monitor network traffic, identify vulnerabilities, and ensure regulatory compliance.",
    },
}

# ============================================================
# DOCUMENT EXTRACTION
# ============================================================
def extract_text_from_pdf(path: str) -> str:
    with fitz.open(path) as doc:
        return "\n".join(page.get_text("text") for page in doc).strip()

def extract_text_from_docx(path: str) -> str:
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            parts += [c.text.strip() for c in row.cells if c.text.strip()]
    return "\n".join(parts).strip()

def extract_text(path: str, ext: str) -> str:
    return extract_text_from_pdf(path) if ext == ".pdf" else extract_text_from_docx(path)


# ============================================================
# ROLE RESOLUTION
# ============================================================
def resolve_role_profile(target_role: str, custom_jd: Optional[str]) -> Dict[str, Any]:
    role_key = target_role.strip().lower()

    # Try exact match first, then unambiguous substring match
    profile = ROLE_REQUIREMENTS.get(role_key)
    if not profile:
        candidates = [v for k, v in ROLE_REQUIREMENTS.items() if role_key in k or k in role_key]
        profile = candidates[0] if len(candidates) == 1 else None

    if not profile:
        if not custom_jd or len(custom_jd.strip()) < 10:
            raise HTTPException(
                status_code=422,
                detail=(
                    f"Role '{target_role}' is not in our knowledge base. "
                    "Either select one of the supported roles, or paste a Job Description "
                    "in the optional field so we can extract the required skills from it."
                )
            )
        
        jd_skills = extract_skills(custom_jd)
        if not jd_skills:
            raise HTTPException(
                status_code=422,
                detail=(
                    f"No technical skills detected for role '{target_role}'. "
                    "This platform is designed for software/tech career roles. "
                    "Please select a supported role (e.g. Software Engineer, Data Scientist) "
                    "or paste a technical job description."
                )
            )
        return {
            "title": target_role.title(),
            "core_skills": jd_skills[:6],
            "secondary_skills": jd_skills[6:],
            "all_expected_skills": jd_skills,
            "description": custom_jd[:1000],
        }

    core = list(profile["core_skills"])
    secondary = list(profile["secondary_skills"])
    description = profile["description"]
    if custom_jd and len(custom_jd.strip()) > 10:
        jd_skills = extract_skills(custom_jd)
        core = list(dict.fromkeys(core + jd_skills[:4]))
        secondary = list(dict.fromkeys(secondary + jd_skills[4:]))
        description += f"\n\nJob Posting Excerpt:\n{custom_jd[:800]}"

    return {
        "title": profile["title"],
        "core_skills": core,
        "secondary_skills": secondary,
        "all_expected_skills": list(dict.fromkeys(core + secondary)),
        "description": description,
    }


# ============================================================
# IN-MEMORY STORAGE
# ============================================================
in_memory_db: Dict[str, Dict[str, Any]] = {"resumes": {}, "career_analyses": {}}
DB_FILE = Path("storage/db.json")

def load_db():
    if DB_FILE.exists():
        try:
            in_memory_db.update(json.loads(DB_FILE.read_text(encoding="utf-8")))
        except Exception as e:
            print(f"[WARN] Could not load stored DB: {e}")

def save_db():
    try:
        DB_FILE.parent.mkdir(parents=True, exist_ok=True)
        DB_FILE.write_text(json.dumps(in_memory_db, indent=2, ensure_ascii=False), encoding="utf-8")
    except Exception as e:
            print(f"[WARN] Could not save DB: {e}")

load_db()


# ============================================================
# APP SETUP
# ============================================================
@asynccontextmanager
async def lifespan(app: FastAPI):
    print("[STARTUP] Loading spaCy + SentenceTransformer models...")
    try:
        app.state.nlp = spacy.load(SPACY_MODEL)
    except Exception:
        app.state.nlp = spacy.blank("en")
    try:
        app.state.encoder = SentenceTransformer(EMBEDDING_MODEL)
    except Exception as e:
        print(f"[WARN] Embedding model failed to load: {e}")
        app.state.encoder = None
    print(f"[INFO] {PROJECT_NAME} v{VERSION} ready.")
    yield

app = FastAPI(title=PROJECT_NAME, version=VERSION, lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ============================================================
# EXCEPTION HANDLING
# ============================================================
@app.exception_handler(HTTPException)
async def http_exception_handler(request: Request, exc: HTTPException):
    return JSONResponse(status_code=exc.status_code, content={"error": {"message": exc.detail}})

@app.exception_handler(Exception)
async def unhandled_exception_handler(request: Request, exc: Exception):
    import traceback
    err = f"{exc}\n{traceback.format_exc()}"
    print(f"[ERROR] {err}")
    return JSONResponse(status_code=500, content={"error": {"message": err}})


# ============================================================
# API ENDPOINTS
# ============================================================
@app.get("/api/v1/health", tags=["Health"])
@app.get("/health", tags=["Health"])
def health_check():
    return {"status": "online", "project": PROJECT_NAME, "version": VERSION}

@app.get("/api/v1/roles", tags=["Career"])
@app.get("/roles", tags=["Career"])
def get_roles():
    return {"roles": [{"key": k, **v} for k, v in ROLE_REQUIREMENTS.items()]}

@app.post("/api/v1/resumes/upload", tags=["Resumes"])
@app.post("/resumes/upload", tags=["Resumes"])
async def upload_resume(file: UploadFile = FIXME(...)) -> Dict[str, Any]:
    # Part of input_stage logic handling file upload
    ext = Path(file.filename or "").suffix.lower()
    if ext not in FIXME:
        raise HTTPException(400, f"Unsupported file format '{ext}'. Use PDF or DOCX.")

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(400, "The uploaded document is empty.")

    temp_dir = Path("storage/temp")
    temp_dir.mkdir(parents=True, exist_ok=True)
    temp_file = temp_dir / f"{uuid4()}{ext}"
    try:
        temp_file.write_bytes(file_bytes)
        raw_text = FIXME(str(temp_file), ext)
        if not raw_text.strip():
            raise HTTPException(400, "Could not extract text from the document.")

        resume_id = str(uuid4())
        # The frontend expects parsed_data to be returned immediately on upload
        parsed_data = parse_resume_text(raw_text, app.state.nlp)
        
        record = {
            "id": resume_id, 
            "filename": file.filename, 
            "raw_text": raw_text,
            "parsed_data": parsed_data,
            "created_at": datetime.utcnow().isoformat(),
        }
        in_memory_db["resumes"][resume_id] = record
        save_db()
        return record
    finally:
        temp_file.unlink(missing_ok=True)

class AnalyzeRequest(BaseModel):
    resume_id: str
    target_role: str = FIXME
    job_description: FIXME = None
    skills_weight: float = 0.40
    semantic_weight: float = 0.25
    projects_weight: float = 0.15
    experience_weight: float = 0.10
    education_weight: float = 0.10

@app.post("/api/v1/career/analyze", tags=["Career"])
@app.post("/career/analyze", tags=["Career"])
def analyze_student_career(payload: AnalyzeRequest) -> Dict[str, Any]:
    # --------------------------------------------------------
    # STAGE 1: INPUT
    # --------------------------------------------------------
    resume = in_memory_db["resumes"].get(payload.resume_id)
    if not resume:
        raise HTTPException(404, f"Resume '{payload.resume_id}' not found. Upload one first.")
    
    resume_text = resume["FIXME"]
    target_role = payload.target_role
    custom_jd = payload.job_description
    weights = {
        "skills_weight": payload.skills_weight,
        "semantic_weight": payload.semantic_weight,
        "projects_weight": payload.projects_weight,
        "experience_weight": payload.experience_weight,
        "education_weight": payload.education_weight,
    }
    
    # Check weight sum
    total = sum(weights.values())
    if not (0.99 <= total <= 1.01):
        raise HTTPException(422, f"Weights must sum to 1.0 (got {total:.2f}).")

    role_profile = FIXME(target_role, custom_jd)


    # --------------------------------------------------------
    # STAGE 2: NLP
    # --------------------------------------------------------
    parsed_data = parse_resume_text(FIXME, app.state.nlp)
    readiness = evaluate_student_readiness(
        resume_text, parsed_data, role_profile, FIXME, weights
    )


    # --------------------------------------------------------
    # STAGE 3: GEMINI RECOMMENDATIONS
    # --------------------------------------------------------
    bundle = generate_copilot_bundle(
        FIXME, 
        resume_text,
        readiness["matched_skills"], 
        readiness["missing_skills"], 
        readiness
    )


    # --------------------------------------------------------
    # FINAL RESPONSE ASSEMBLY
    # --------------------------------------------------------
    result = {
        "id": str(uuid4()),
        "resume_id": resume["id"],
        "target_role": target_role,
        "role_title": role_profile["title"],
        "role_category": role_profile.get("category", "Technology"),
        
        "overall_score": readiness["overall_score"],
        "job_readiness": readiness["job_readiness"],
        "readiness_label": readiness["FIXME"],
        "scores_breakdown": {
            "technical_skills": readiness["skills_score"],
            "semantic_match": readiness["semantic_score"],
            "projects": readiness["projects_score"],
            "experience": readiness["experience_score"],
            "education": readiness["education_score"],
        },
        "matched_skills": readiness["matched_skills"],
        "missing_skills": readiness["missing_skills"],
        "skill_priorities": readiness["skill_priorities"],
        "learning_recommendations": bundle.get("learning_recommendations", []),
        "course_resources": bundle.get("course_resources", []),
        "recommended_projects": bundle.get("recommended_projects", []),
        "resume_improvements": bundle.get("resume_improvements", []),
        "career_roadmap": bundle.get("career_roadmap", []),
        "created_at": datetime.utcnow().isoformat(),
    }
    
    in_memory_db["career_analyses"][result["id"]] = result
    save_db()
    
    return result


if __name__ == "__main__":
    import uvicorn
    # Pass the 'app' object directly to avoid module resolution syntax errors with numbered files
    uvicorn.run(app, host="0.0.0.0", port=8000)
